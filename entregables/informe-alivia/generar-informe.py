"""Reconstrucción: python generar-informe.py --root RUTA_REPOSITORIO [--salida DIRECTORIO].
Requiere Python 3, python-docx y Pillow; exportación e índices: Microsoft Word de escritorio
y PowerShell 7 en Windows. Usa originales del repositorio en modo lectura. No instala paquetes.
"""
from pathlib import Path
import argparse,subprocess,tempfile,re,sys,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from contenido import TITLE,NAMES,ABSTRACT,CHAPTERS,REFERENCES,APPENDIX,EVIDENCE
from figuras import build

def xml(parent,tag,attrs=None):
 el=OxmlElement(tag)
 for key,value in (attrs or {}).items():el.set(qn(key),str(value))
 parent.append(el);return el
def field(p,instruction):
 run=p.add_run();xml(run._r,'w:fldChar',{'w:fldCharType':'begin'})
 r=p.add_run();t=xml(r._r,'w:instrText',{'xml:space':'preserve'});t.text=' '+instruction+' '
 xml(p.add_run()._r,'w:fldChar',{'w:fldCharType':'separate'})
 p.add_run('0')
 xml(p.add_run()._r,'w:fldChar',{'w:fldCharType':'end'})
def rich(p,text):
 for i,piece in enumerate(re.split(r'\*([^*]+)\*',text)):
  p.add_run(piece).italic=bool(i%2)
def set_font(style,size=12,bold=False,italic=False):
 style.font.name='Times New Roman';style.font.size=Pt(size);style.font.bold=bold;style.font.italic=italic;style.font.color.rgb=RGBColor(0,0,0)
 rpr=style.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
 if fonts is None:fonts=xml(rpr,'w:rFonts')
 for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  if qn('w:'+attr) in fonts.attrib:del fonts.attrib[qn('w:'+attr)]
def make_style(doc,name,size=12,bold=False,italic=False):
 style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 style.base_style=doc.styles['Normal'];set_font(style,size,bold,italic)
 pf=style.paragraph_format;pf.first_line_indent=Inches(0);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=2
 return style
def section_setup(sec,landscape=False):
 sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
 sec.page_width=Inches(11 if landscape else 8.5);sec.page_height=Inches(8.5 if landscape else 11)
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
 sec.header_distance=sec.footer_distance=Inches(.5)
def main():
 parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[2]);parser.add_argument('--salida',type=Path,default=Path(__file__).resolve().parent);args=parser.parse_args()
 root=args.root.resolve();out=args.salida.resolve();out.mkdir(parents=True,exist_ok=True)
 assert (root/'docs/07-datos/00-diagrama-er-global.md').is_file()
 assert 150<=len(ABSTRACT.split())<=220
 build(out/'recursos',root)
 doc=Document();section_setup(doc.sections[0]);doc.core_properties.title=TITLE;doc.core_properties.subject='Ingeniería de Software';doc.core_properties.author='; '.join(NAMES)
 normal=doc.styles['Normal'];set_font(normal);normal.paragraph_format.line_spacing=2;normal.paragraph_format.first_line_indent=Inches(.5);normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.widow_control=True
 lang=xml(doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr')),'w:lang',{'w:val':'es-PE'})
 for level in range(1,4):
  st=doc.styles['Heading '+str(level)];set_font(st,bold=True,italic=level==3);st.paragraph_format.first_line_indent=Inches(0);st.paragraph_format.line_spacing=2;st.paragraph_format.space_before=st.paragraph_format.space_after=Pt(0);st.paragraph_format.keep_with_next=True;st.paragraph_format.keep_together=True;st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
  st.paragraph_format.page_break_before=level==1
 title=make_style(doc,'Title',12,True);title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
 front=make_style(doc,'Front Heading',12,True);front.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;front.paragraph_format.page_break_before=True;front.paragraph_format.keep_with_next=True
 for name in ['Figure Caption','Table Caption']:
  st=make_style(doc,name);st.paragraph_format.keep_with_next=True;st.paragraph_format.keep_together=True
 note=make_style(doc,'Figure Note',10);note.paragraph_format.line_spacing=1.0;note.paragraph_format.space_after=Pt(6)
 imgstyle=make_style(doc,'Figure Image');imgstyle.paragraph_format.line_spacing=1;imgstyle.paragraph_format.keep_with_next=True;imgstyle.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for n in range(1,10):
  st=make_style(doc,'TOC '+str(n));st.paragraph_format.line_spacing=1.15;st.paragraph_format.keep_with_next=False;st.paragraph_format.left_indent=Inches(.25*(n-1));st.paragraph_format.right_indent=Inches(0)
 reference=make_style(doc,'Reference');reference.paragraph_format.left_indent=Inches(.5);reference.paragraph_format.first_line_indent=Inches(-.5)
 header=doc.sections[0].header.paragraphs[0];header.style=doc.styles['Normal'];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT;header.paragraph_format.first_line_indent=Inches(0);header.paragraph_format.line_spacing=1;field(header,'PAGE')
 doc.settings.element.append(OxmlElement('w:updateFields'));doc.settings.element[-1].set(qn('w:val'),'true')
 # La portada utiliza espaciado de párrafo, no renglones vacíos para simular posiciones.
 cover=doc.add_paragraph(style='Title');cover.paragraph_format.space_before=Pt(72)
 cover.add_run('Alivía: Análisis y diseño de un sistema');cover.add_run().add_break();cover.add_run('de coordinación y trazabilidad de ayuda humanitaria')
 for i,name in enumerate(NAMES):
  pa=doc.add_paragraph(name);pa.paragraph_format.first_line_indent=Inches(0);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
  if i==0:pa.paragraph_format.space_before=Pt(24)
 for text in ['Universidad Nacional de Piura (UNP)','Curso: Ingeniería de Software','13 de octubre de 2026']:
  pa=doc.add_paragraph(text);pa.paragraph_format.first_line_indent=Inches(0);pa.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.add_paragraph('Resumen',style='Front Heading');pa=doc.add_paragraph(ABSTRACT);pa.paragraph_format.first_line_indent=Inches(0)
 pa=doc.add_paragraph();pa.paragraph_format.first_line_indent=Inches(.5);r=pa.add_run('Palabras clave: ');r.italic=True;pa.add_run('ayuda humanitaria, trazabilidad, inventario')
 for title,instruction in [('Índice general','TOC \\o "1-2" \\h \\z \\u'),('Índice de figuras','TOC \\h \\z \\c "Figura"'),('Índice de tablas','TOC \\h \\z \\c "Tabla"')]:
  doc.add_paragraph(title,style='Front Heading');field(doc.add_paragraph(),' '+instruction)
 count={'fig':0,'table':0};captions=[]
 def add_caption(kind,title,annex=False):
  label='Figura' if kind=='fig' else 'Tabla';seq=label
  count[kind]+=1;pa=doc.add_paragraph(style='Figure Caption' if kind=='fig' else 'Table Caption')
  pa.add_run(label+(' A' if annex else ' ')).bold=True;field(pa,'SEQ '+seq+(' \\r 1' if annex else '')+' \\* ARABIC')
  for run in pa.runs:run.bold=True
  pa.add_run(' ').bold=True
  pa.add_run().add_break();pa.add_run(title).italic=True
  captions.append({'kind':kind,'title':title,'annex':annex})
 def block(b,annex=False):
  kind=b['type']
  if kind=='p':
   text=b['text']
   if text=='EVIDENCIA_LOCAL':
    # La reconstrucción conserva el resultado observado; no ejecuta ni sobrescribe prototipos.
    text=EVIDENCE
   rich(doc.add_paragraph(),text)
  elif kind=='h2':doc.add_heading(b['text'],2)
  elif kind=='fig':
   if b['landscape']:section_setup(doc.add_section(WD_SECTION_START.NEW_PAGE),True)
   add_caption('fig',b['title']);path=out/'recursos'/(b['key']+'.png');im=Image.open(path);maxw=9 if b['landscape'] else 6.5;maxh=4.84 if b['landscape'] else 5.25
   w=min(maxw,maxh*im.width/im.height)
   pa=doc.add_paragraph(style='Figure Image');pa.add_run().add_picture(str(path),width=Inches(w))
   pa=doc.add_paragraph(style='Figure Note');pa.add_run('Nota. ').italic=True;rich(pa,b['note'])
   if b['landscape']:section_setup(doc.add_section(WD_SECTION_START.NEW_PAGE),False)
  elif kind=='table':
   add_caption('table',b['title'],annex);t=doc.add_table(rows=1,cols=len(b['headers']));t.autofit=False
   widths=[2.15,4.35]
   for col,width in zip(t.columns,widths):col.width=Inches(width)
   tPr=t._tbl.tblPr;borders=xml(tPr,'w:tblBorders')
   for side in ['top','bottom']:xml(borders,'w:'+side,{'w:val':'single','w:sz':'6','w:color':'000000'})
   for side in ['left','right','insideV','insideH']:xml(borders,'w:'+side,{'w:val':'nil'})
   for j,val in enumerate(b['headers']):t.rows[0].cells[j].text=val
   xml(t.rows[0]._tr.get_or_add_trPr(),'w:tblHeader')
   for row in b['rows']:
    cells=t.add_row().cells
    for cell,text in zip(cells,row):cell.text=text
   for i,row in enumerate(t.rows):
    xml(row._tr.get_or_add_trPr(),'w:cantSplit')
    for j,cell in enumerate(row.cells):
     cell.width=Inches(widths[j])
     if i==0:
      cb=xml(cell._tc.get_or_add_tcPr(),'w:tcBorders');xml(cb,'w:bottom',{'w:val':'single','w:sz':'6'})
     for pa in cell.paragraphs:
      pa.paragraph_format.first_line_indent=Inches(0);pa.paragraph_format.line_spacing=1;pa.paragraph_format.space_before=Pt(3);pa.paragraph_format.space_after=Pt(3);pa.paragraph_format.keep_with_next=i<len(t.rows)-1
      for run in pa.runs:run.font.name='Times New Roman';run.font.size=Pt(11);run.bold=i==0
   if b['note']:
    pa=doc.add_paragraph(style='Figure Note');pa.add_run('Nota. ').italic=True;rich(pa,b['note'])
 for title,blocks in CHAPTERS:
  doc.add_heading(title,1)
  for b in blocks:block(b)
 doc.add_heading('Referencias',1)
 for ref in REFERENCES:rich(doc.add_paragraph(style='Reference'),ref)
 doc.add_heading('Anexo A. Trazabilidad de requisitos y verificaciones',1)
 for b in APPENDIX:block(b,True)
 assert count=={'fig':18,'table':11},count
 for parent in [doc.styles.element,doc.element]:
  for border in list(parent.iter(qn('w:pBdr'))):border.getparent().remove(border)
 target=out/'Alivia_Informe_Analisis_Diseno.docx';doc.save(target)
 result=subprocess.run(['pwsh.exe','-NoProfile','-File',str(Path(__file__).parent/'actualizar-exportar.ps1'),'-InputDocx',str(target),'-OutputPdf',str(out/'Alivia_Informe_Analisis_Diseno.pdf')],capture_output=True,text=True,encoding='utf-8')
 if result.returncode:raise RuntimeError(result.stderr+'\n'+result.stdout)
 print(result.stdout)
 print(json.dumps({'figuras':count['fig'],'tablas':count['table'],'resumen_palabras':len(ABSTRACT.split()),'docx':str(target)},ensure_ascii=False))
if __name__=='__main__':
 sys.stdout.reconfigure(encoding='utf-8');main()
